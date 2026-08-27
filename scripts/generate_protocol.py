# -*- coding: utf-8 -*-
"""Генерация проекта протокола разногласий."""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from docx_helpers import (
    add_paragraph,
    repeat_header_row,
    set_run_font,
    set_table_width,
    setup_page,
    write_cell,
)
from protocol_rows import ROWS


OUT = Path("/workspace/docs/Protokol_raznoglasiy.docx")
TABLE_WIDTH = 9922  # twips ≈ 17.5 cm


def build():
    document = Document()
    setup_page(document)

    add_paragraph(
        document,
        "Протокол разногласий",
        size=16,
        bold=True,
        align="center",
        space_after=0,
    )
    add_paragraph(
        document,
        "к проекту договора о передаче в безвозмездное пользование",
        size=13,
        bold=True,
        align="center",
        space_after=0,
    )
    add_paragraph(
        document,
        "муниципального имущества",
        size=13,
        bold=True,
        align="center",
        space_after=12,
    )

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run1 = p.add_run("г. Чебоксары")
    set_run_font(run1, size=12, bold=True)
    run2 = p.add_run("\t\t\t\t«___» _____________ 2026 г.")
    set_run_font(run2, size=12)

    add_paragraph(
        document,
        "«___» _____________ 2026 г. муниципальным автономным образовательным учреждением "
        "дополнительного образования «Дворец детского (юношеского) творчества» города Чебоксары "
        "Чувашской Республики (МАОУДО «ДДЮТ» г. Чебоксары) получен проект договора о передаче "
        "в безвозмездное пользование муниципального имущества.",
        first_line=1.25,
    )
    add_paragraph(
        document,
        "По этому договору МАОУДО «ДДЮТ» г. Чебоксары выступает Ссудодателем, а Автономная "
        "профессиональная образовательная некоммерческая организация «Сингулярити Хаб» "
        "(Центр Сингулярности) (АПОНО «Сингулярити Хаб») в лице директора Соловьева Георгия "
        "Михайловича, действующего на основании Устава, — Ссудополучателем.",
        first_line=1.25,
    )
    add_paragraph(
        document,
        "МАОУДО «ДДЮТ» г. Чебоксары не согласно со следующими условиями проекта договора и "
        "предлагает согласовать их в следующей редакции:",
        first_line=1.25,
        space_after=10,
    )

    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, TABLE_WIDTH)

    widths = [Cm(2.2), Cm(5.1), Cm(5.1), Cm(5.1)]
    for i, width in enumerate(widths):
        for cell in table.columns[i].cells:
            cell.width = width

    headers = [
        "№ пункта договора",
        "Редакция Ссудополучателя",
        "Редакция Ссудодателя",
        "Обоснование редакции Ссудодателя",
    ]
    hdr = table.rows[0]
    repeat_header_row(hdr)
    for i, text in enumerate(headers):
        write_cell(hdr.cells[i], text, size=10, bold=True, align="center", shade="1F4E79")
        for p in hdr.cells[i].paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)

    for row_data in ROWS:
        row = table.add_row()
        write_cell(row.cells[0], row_data["punkt"], size=9, bold=True, align="center")
        write_cell(row.cells[1], row_data["other"], size=8, align="justify")
        write_cell(row.cells[2], row_data["ours"], size=8, align="justify")
        write_cell(row.cells[3], row_data["reason"], size=8, align="justify")

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    add_paragraph(document, "", space_after=6)

    points = (
        "преамбулы, правил наименования предмета по тексту договора, "
        "п. п. 1.1, 1.2, 1.3, 1.4, 1.6, 1.7, 1.8, 1.9, 2.1, 2.2, 2.3, 2.4, "
        "3.1.1, 3.1.3, 3.1.5, 3.1.6, 3.1.7, 3.1.9, 3.1.10, 3.1.11, 3.1.12, 3.1.13, "
        "3.1.17, 3.2, 3.3.1, 3.3.4, 3.3.5, 3.3.6, 3.3.10 (излагается как п. 3.3.9), "
        "3.3.14, 3.3.21, 3.4.1, 3.4.3, 3.4.4, 4.1, 4.2, 5.1, 5.2, 6.2, 6.3, 6.4, 6.5, "
        "7.4 договора и приложения № 1"
    )

    add_paragraph(
        document,
        f"Подписывая настоящий протокол, МАОУДО «ДДЮТ» г. Чебоксары и АПОНО «Сингулярити Хаб» "
        f"подтверждают, что условия {points} будут действовать в редакции Ссудодателя, изложенной "
        f"в настоящем протоколе. Условия, не указанные в протоколе, действуют в редакции проекта "
        f"договора.",
        first_line=1.25,
    )
    add_paragraph(
        document,
        "Ссудодатель предлагает Ссудополучателю направить письменный ответ о принятии или "
        "отклонении условий в редакции Ссудодателя в течение семи дней со дня получения "
        "настоящего протокола. Неполучение ответа в указанный срок не считается акцептом "
        "редакции Ссудодателя.",
        first_line=1.25,
    )
    add_paragraph(
        document,
        "Протокол составлен в двух экземплярах, имеющих одинаковую юридическую силу, по одному "
        "для каждой из Сторон, и после подписания является неотъемлемой частью договора.",
        first_line=1.25,
        space_after=14,
    )

    add_paragraph(
        document,
        "Адреса и реквизиты сторон",
        size=13,
        bold=True,
        align="center",
        space_after=10,
    )

    parties = document.add_table(rows=1, cols=2)
    parties.autofit = True
    set_table_width(parties, TABLE_WIDTH)
    left, right = parties.rows[0].cells
    left.width = Cm(8.75)
    right.width = Cm(8.75)

    def fill_party(cell, lines):
        cell.text = ""
        for i, (text, bold) in enumerate(lines):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(text)
            set_run_font(run, size=11, bold=bold)

    fill_party(
        left,
        [
            ("Ссудодатель:", True),
            ("Муниципальное автономное образовательное учреждение дополнительного образования "
             "«Дворец детского (юношеского) творчества» города Чебоксары Чувашской Республики", False),
            ("(МАОУДО «ДДЮТ» г. Чебоксары)", False),
            ("Адрес: Чувашская Республика, город Чебоксары, Президентский бульвар, д. 14", False),
            ("ОГРН: 1022101136155", False),
            ("ИНН: 2128024030  КПП: 213001001", False),
            ("", False),
            ("Директор", False),
            ("", False),
            ("_______________ / Воробьева Е.В. /", False),
            ("М.П.", False),
        ],
    )
    fill_party(
        right,
        [
            ("Ссудополучатель:", True),
            ("Автономная профессиональная образовательная некоммерческая организация "
             "«Сингулярити Хаб» (Центр Сингулярности)", False),
            ("(АПОНО «Сингулярити Хаб»)", False),
            ("Адрес: 428003, Чувашская Республика, г.о. город Чебоксары, г. Чебоксары, "
             "ул. К. Маркса, д. 47, помещ. 9", False),
            ("ОГРН: 1242100004640", False),
            ("ИНН: 2100017604  КПП: 210001001", False),
            ("Эл. почта: it.singularityhub@yandex.ru", False),
            ("Директор", False),
            ("", False),
            ("_______________ / Соловьев Г.М. /", False),
            ("М.П.", False),
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUT)
    print(f"saved {OUT}")


if __name__ == "__main__":
    build()
